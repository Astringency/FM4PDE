from __future__ import annotations

import csv
import itertools
import json
from pathlib import Path

from openpyxl import load_workbook

from scripts.summary import MAIN_COLUMNS, summary
from scripts.analysis.summarize_poisson_sampler_comparison import collect_comparison


def _write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def _main_row(*, n: int, key: str, pde: str = "poisson") -> dict[str, object]:
    return {
        "ablation_group_key": key,
        "pde": pde,
        "task": "both",
        "sensor_mode": "sensor_column",
        "rel_l2_a_mean": 0.1,
        "rel_l2_u_mean": 0.2,
        "pde_residual_norm_mean": 9.0,
        "obs_rel_l2_a_mean": 8.0,
        "obs_rel_l2_u_mean": 7.0,
        "rel_l2_a_n": n,
    }


def _write_main_experiment(path: Path, *, dist_key: str) -> None:
    sample_key = f"pde=poisson|resolved_residual_mode=|dist={dist_key}"
    run_key = f"pde=poisson|resolved_residual_mode=static|dist={dist_key}"
    _write_csv(
        path / "summary_all_grouped.csv",
        [
            _main_row(n=1000, key=sample_key),
            _main_row(n=4, key=f"{sample_key}|pilot=true", pde="pilot"),
        ],
    )
    _write_csv(
        path / "summary_run_seed_grouped.csv",
        [
            {
                "ablation_group_key": run_key,
                "L_pde_mean": 0.3,
                "L_obs_a_mean": 0.4,
                "L_obs_u_mean": 0.5,
            }
        ],
    )


def test_main_summary_has_exact_sheets_columns_and_complete_groups(tmp_path: Path) -> None:
    outputs = tmp_path / "outputs"
    _write_main_experiment(outputs / "main" / "MAIN1000_100_TEST_id", dist_key="id")
    _write_main_experiment(
        outputs / "main" / "MAIN1000_100_TEST_FULL_rough", dist_key="rough"
    )

    written = summary(outputs, "main")

    assert set(written) == {"main"}
    workbook = load_workbook(written["main"], read_only=True, data_only=True)
    assert workbook.sheetnames == ["FM4PDE_SPARSE", "FM4PDE_FULL"]
    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))
        assert rows[0] == MAIN_COLUMNS
        assert len(rows) == 2
        assert rows[1][3] == "sensor_col"
        assert rows[1][4:9] == (0.1, 0.2, 0.3, 0.4, 0.5)


def test_ablation_summary_groups_pdes_by_ablation_type(tmp_path: Path) -> None:
    outputs = tmp_path / "outputs"
    for pde, noise in (("burger", 0.0), ("poisson", 0.1)):
        run_dir = (
            outputs
            / "ablations"
            / pde
            / "both"
            / f"noise_robustness_{pde}_000"
            / "20260903-000000"
        )
        run_dir.mkdir(parents=True)
        (run_dir / "resolved_config.yaml").write_text(
            "\n".join(
                (
                    f"pde: {pde}",
                    "task: both",
                    "test_type: smooth",
                    "ablation_group: noise_robustness",
                    f"ablation_name: noise_robustness_{pde}_000",
                    f"noise_level: {noise}",
                )
            )
            + "\n",
            encoding="utf-8",
        )
        (run_dir / "metrics_final.json").write_text(
            json.dumps(
                {
                    "status": "ok",
                    "pde_residual_status": "reliable",
                    "rel_l2_a": 0.1,
                    "rel_l2_u": 0.2,
                    "L_pde": 0.3,
                    "L_obs_a": 0.4,
                    "L_obs_u": 0.5,
                }
            ),
            encoding="utf-8",
        )

    written = summary(outputs, "ablations")

    assert set(written) == {"ablations"}
    workbook = load_workbook(written["ablations"], read_only=True, data_only=True)
    assert workbook.sheetnames == ["noise_robustness"]
    rows = list(workbook["noise_robustness"].iter_rows(values_only=True))
    assert rows[0] == (
        "PDE",
        "TASK",
        "DIST",
        "NOISE",
        "rel L2(a)",
        "rel L2(u)",
        "pde L",
        "obs L(a)",
        "obs L(u)",
        "Remark",
    )
    assert [row[0] for row in rows[1:]] == ["burger", "poisson"]
    assert {row[2] for row in rows[1:]} == {"Smooth"}


def test_poisson_sampler_comparison_requires_and_returns_all_36_runs(tmp_path: Path) -> None:
    root = tmp_path / "comparison"
    phases = ("stochastic", "deterministic", "hybrid_d2s", "hybrid_s2d")
    for dist, task, phase in itertools.product(
        ("id", "smooth", "rough"), ("both", "forward", "inverse"), phases
    ):
        run = root / dist / task / phase
        run.mkdir(parents=True)
        (run / "resolved_config.yaml").write_text(
            "\n".join(
                (
                    "pde: poisson",
                    f"task: {task}",
                    f"test_type: {dist}",
                    f"data_path: /data/poisson_{dist}.mat",
                    "ablation_group: poisson_sampler_comparison",
                    f"ablation_name: poisson_sampler_comparison_poisson_{phases.index(phase):03d}",
                    f"sampler_phase: {phase}",
                    "switch_ratio: 0.5",
                    "sample_seed: 0",
                    "offset: 0",
                    "batch_size: 1",
                )
            )
            + "\n",
            encoding="utf-8",
        )
        (run / "metrics_final.json").write_text(
            json.dumps(
                {
                    "status": "ok",
                    "pde_residual_status": "reliable",
                    "rel_l2_a": 0.1,
                    "rel_l2_u": 0.2,
                    "L_pde": 0.3,
                    "L_obs_a": 0.4,
                    "L_obs_u": 0.5,
                }
            ),
            encoding="utf-8",
        )

    rows = collect_comparison(root)

    assert len(rows) == 36
    assert {row["DIST"] for row in rows} == {"ID", "Smooth", "Rough"}
    assert {row["SAMPLER"] for row in rows} == set(phases)
