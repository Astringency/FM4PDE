#!/usr/bin/env python3
"""Export MAIN1000 results after row-level tuned1 replacement to Excel."""

from __future__ import annotations

import argparse
import csv
import math
import statistics
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


TEST_ORDER = {"id": 0, "rough": 1, "smooth": 2}
PDE_ORDER = {
    "burger": 0,
    "poisson": 1,
    "helmholtz": 2,
    "darcy": 3,
    "nsnonbounded": 4,
}
TASK_ORDER = {"forward": 0, "inverse": 1, "both": 2}

INK = "172033"
BLUE = "185FA5"
BLUE_DARK = "0B4778"
BLUE_LIGHT = "DCEAF7"
BLUE_XLIGHT = "EEF5FB"
GOLD_LIGHT = "FFF3CD"
GREY = "667085"
GREY_LIGHT = "E9EDF2"
WHITE = "FFFFFF"


def _as_float(value: str | None) -> float | None:
    if value is None or value == "":
        return None
    return float(value)


def _as_int(value: str | None) -> int | None:
    parsed = _as_float(value)
    return None if parsed is None else int(parsed)


def _as_bool(value: str | None) -> bool:
    return str(value).strip().lower() in {"true", "1", "yes"}


def _read_comparisons(path: Path) -> list[dict[str, Any]]:
    with path.open(newline="", encoding="utf-8") as handle:
        raw_rows = list(csv.DictReader(handle))

    rows: list[dict[str, Any]] = []
    for raw in raw_rows:
        row: dict[str, Any] = {
            "test_type": raw["test_type"],
            "pde": raw["pde"],
            "task": raw["task"],
            "sensor_mode": raw["sensor_mode"],
            "primary_metric": raw["primary_metric"],
            "baseline_n": _as_int(raw["baseline_n"]),
            "tuned1_available": _as_bool(raw["tuned1_available"]),
            "tuned1_n": _as_int(raw["tuned1_n"]),
        }
        for field in (
            "baseline_rel_l2_a_mean",
            "baseline_rel_l2_a_p90",
            "baseline_rel_l2_u_mean",
            "baseline_rel_l2_u_p90",
            "baseline_pde_residual_mean",
            "baseline_primary_mean",
            "tuned1_rel_l2_a_mean",
            "tuned1_rel_l2_a_p90",
            "tuned1_rel_l2_u_mean",
            "tuned1_rel_l2_u_p90",
            "tuned1_pde_residual_mean",
            "tuned1_primary_mean",
            "primary_improvement",
            "rel_l2_a_improvement",
            "rel_l2_u_improvement",
            "pde_residual_improvement",
        ):
            row[field] = _as_float(raw[field])
        rows.append(row)

    rows.sort(
        key=lambda row: (
            TEST_ORDER[row["test_type"]],
            PDE_ORDER[row["pde"]],
            TASK_ORDER[row["task"]],
            row["sensor_mode"],
        )
    )
    return rows


def _select_best(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    selected: list[dict[str, Any]] = []
    for source in rows:
        row = dict(source)
        base_primary = row["baseline_primary_mean"]
        tuned_primary = row["tuned1_primary_mean"]
        use_tuned = bool(
            row["tuned1_available"]
            and tuned_primary is not None
            and base_primary is not None
            and tuned_primary < base_primary
        )
        prefix = "tuned1" if use_tuned else "baseline"
        row.update(
            {
                "result_source": prefix,
                "replaced": use_tuned,
                "selection_reason": (
                    "tuned1 任务主指标更低"
                    if use_tuned
                    else ("无 tuned1 结果" if not row["tuned1_available"] else "tuned1 未改善主指标")
                ),
                "final_n": row[f"{prefix}_n"],
                "final_primary_mean": row[f"{prefix}_primary_mean"],
                "final_rel_l2_a_mean": row[f"{prefix}_rel_l2_a_mean"],
                "final_rel_l2_a_p90": row[f"{prefix}_rel_l2_a_p90"],
                "final_rel_l2_u_mean": row[f"{prefix}_rel_l2_u_mean"],
                "final_rel_l2_u_p90": row[f"{prefix}_rel_l2_u_p90"],
                "final_pde_residual_mean": row[f"{prefix}_pde_residual_mean"],
            }
        )
        selected.append(row)
    return selected


def _validate_rows(rows: list[dict[str, Any]]) -> None:
    keys = [
        (row["test_type"], row["pde"], row["task"], row["sensor_mode"])
        for row in rows
    ]
    assert len(rows) == 42, f"expected 42 result cells, got {len(rows)}"
    assert len(set(keys)) == len(keys), "duplicate TEST/PDE/task/sensor key"
    assert all(row["baseline_n"] == 1000 for row in rows)
    assert all(
        sum(row["test_type"] == test_type for row in rows) == 14
        for test_type in TEST_ORDER
    )
    tuned_rows = [row for row in rows if row["tuned1_available"]]
    assert len(tuned_rows) == 18, f"expected 18 tuned1 cells, got {len(tuned_rows)}"
    assert all(row["tuned1_n"] == 1000 for row in tuned_rows)
    assert all(row["replaced"] for row in tuned_rows)
    for row in tuned_rows:
        expected = (row["baseline_primary_mean"] - row["tuned1_primary_mean"]) / row[
            "baseline_primary_mean"
        ]
        assert math.isclose(expected, row["primary_improvement"], rel_tol=1e-12)


def _style_title(ws, title: str, subtitle: str, end_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    title_cell = ws.cell(1, 1, title)
    title_cell.font = Font(name="Aptos Display", size=20, bold=True, color=INK)
    title_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    subtitle_cell = ws.cell(2, 1, subtitle)
    subtitle_cell.font = Font(name="Aptos", size=10, color=GREY)
    subtitle_cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 32


def _style_header(ws, row_number: int, start_col: int, end_col: int) -> None:
    thin = Side(style="thin", color="CDD5DF")
    for cell in ws.iter_cols(
        min_col=start_col, max_col=end_col, min_row=row_number, max_row=row_number
    ):
        target = cell[0]
        target.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        target.fill = PatternFill("solid", fgColor=BLUE_DARK)
        target.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        target.border = Border(bottom=thin)
    ws.row_dimensions[row_number].height = 30


def _apply_table_style(
    ws,
    header_row: int,
    end_row: int,
    end_col: int,
    source_col: int | None = None,
) -> None:
    _style_header(ws, header_row, 1, end_col)
    thin = Side(style="thin", color="E4E7EC")
    for row_number in range(header_row + 1, end_row + 1):
        for col_number in range(1, end_col + 1):
            cell = ws.cell(row_number, col_number)
            cell.font = Font(name="Aptos", size=10, color=INK)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="center")
            if row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")
        if source_col is not None and ws.cell(row_number, source_col).value == "tuned1":
            for col_number in range(1, end_col + 1):
                ws.cell(row_number, col_number).fill = PatternFill("solid", fgColor=BLUE_XLIGHT)
            ws.cell(row_number, source_col).font = Font(
                name="Aptos", size=10, bold=True, color=BLUE_DARK
            )
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(end_col)}{end_row}"
    ws.freeze_panes = f"A{header_row + 1}"
    ws.sheet_view.showGridLines = False


def _set_widths(ws, widths: dict[int, float]) -> None:
    for col_number, width in widths.items():
        ws.column_dimensions[get_column_letter(col_number)].width = width


def _write_summary(wb: Workbook, rows: list[dict[str, Any]], generated_at: str) -> None:
    ws = wb.active
    ws.title = "摘要"
    _style_title(
        ws,
        "MAIN1000 调参后较优结果汇总",
        "按任务主指标逐行选择较优结果；替换发生时，该行所有最终指标均来自同一个 tuned1 配置。",
        13,
    )

    replacements = [row for row in rows if row["replaced"]]
    residual_improved = [
        row for row in replacements if (row["pde_residual_improvement"] or 0) > 0
    ]
    metrics = [
        ("结果单元", len(rows), "TEST × PDE × task × sensor"),
        ("采用 tuned1", len(replacements), "任务主指标严格低于 baseline"),
        ("保留 baseline", len(rows) - len(replacements), "无 tuned1 或 tuned1 未改善"),
        ("改善中位数", statistics.median(row["primary_improvement"] for row in replacements), "仅 18 个替换项"),
        ("Residual 同时改善", len(residual_improved), "其余替换项 residual 上升"),
    ]
    ws.append([])
    for col, (label, value, note) in enumerate(metrics, start=1):
        cell = ws.cell(4, col * 2 - 1, label)
        cell.font = Font(name="Aptos", size=9, bold=True, color=GREY)
        value_cell = ws.cell(5, col * 2 - 1, value)
        value_cell.font = Font(name="Aptos Display", size=18, bold=True, color=BLUE_DARK)
        note_cell = ws.cell(6, col * 2 - 1, note)
        note_cell.font = Font(name="Aptos", size=8, color=GREY)
        note_cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=4, start_column=col * 2 - 1, end_row=4, end_column=col * 2)
        ws.merge_cells(start_row=5, start_column=col * 2 - 1, end_row=5, end_column=col * 2)
        ws.merge_cells(start_row=6, start_column=col * 2 - 1, end_row=6, end_column=col * 2)
    ws.cell(5, 7).number_format = "0.0%"

    summary_start = 9
    ws.cell(summary_start, 1, "TEST")
    ws.cell(summary_start, 2, "结果数")
    ws.cell(summary_start, 3, "替换数")
    ws.cell(summary_start, 4, "改善中位数")
    ws.cell(summary_start, 5, "改善最小值")
    ws.cell(summary_start, 6, "改善最大值")
    for index, test_type in enumerate(TEST_ORDER, start=summary_start + 1):
        test_rows = [row for row in rows if row["test_type"] == test_type]
        test_replacements = [row for row in test_rows if row["replaced"]]
        improvements = [row["primary_improvement"] for row in test_replacements]
        values = (
            test_type.upper(),
            len(test_rows),
            len(test_replacements),
            statistics.median(improvements),
            min(improvements),
            max(improvements),
        )
        for col_number, value in enumerate(values, start=1):
            ws.cell(index, col_number, value)
    _apply_table_style(ws, summary_start, summary_start + 3, 6)
    for row_number in range(summary_start + 1, summary_start + 4):
        for col_number in (4, 5, 6):
            ws.cell(row_number, col_number).number_format = "0.0%"

    detail_start = 15
    headers = [
        "条目",
        "TEST",
        "PDE",
        "task",
        "主指标改善",
        "Base primary",
        "Tuned1 primary",
        "Residual 改善",
    ]
    for col_number, header in enumerate(headers, start=1):
        ws.cell(detail_start, col_number, header)
    sorted_replacements = sorted(replacements, key=lambda row: row["primary_improvement"])
    for row_number, row in enumerate(sorted_replacements, start=detail_start + 1):
        values = [
            f'{row["test_type"]} | {row["pde"]} | {row["task"]}',
            row["test_type"].upper(),
            row["pde"],
            row["task"],
            row["primary_improvement"],
            row["baseline_primary_mean"],
            row["tuned1_primary_mean"],
            row["pde_residual_improvement"],
        ]
        for col_number, value in enumerate(values, start=1):
            ws.cell(row_number, col_number, value)
    detail_end = detail_start + len(sorted_replacements)
    _apply_table_style(ws, detail_start, detail_end, len(headers))
    for row_number in range(detail_start + 1, detail_end + 1):
        ws.cell(row_number, 5).number_format = "0.0%"
        ws.cell(row_number, 8).number_format = "0.0%"
        for col_number in (6, 7):
            ws.cell(row_number, col_number).number_format = "0.000000"
    ws.conditional_formatting.add(
        f"E{detail_start + 1}:E{detail_end}",
        ColorScaleRule(
            start_type="min",
            start_color=BLUE_LIGHT,
            end_type="max",
            end_color=BLUE_DARK,
        ),
    )

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "tuned1 任务主指标相对改善"
    chart.x_axis.title = "相对改善率"
    chart.x_axis.scaling.min = 0
    chart.x_axis.scaling.max = 0.6
    chart.x_axis.numFmt = "0%"
    chart.y_axis.title = "TEST | PDE | task"
    chart.height = 11
    chart.width = 18
    chart.gapWidth = 45
    chart.legend = None
    chart.dLbls = DataLabelList()
    chart.dLbls.showVal = True
    chart.dLbls.numFmt = "0.0%"
    data = Reference(ws, min_col=5, min_row=detail_start, max_row=detail_end)
    categories = Reference(ws, min_col=1, min_row=detail_start + 1, max_row=detail_end)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.series[0].graphicalProperties.solidFill = BLUE
    chart.series[0].graphicalProperties.line.solidFill = BLUE_DARK
    ws.add_chart(chart, "J9")
    ws.cell(8, 10, "正值表示 tuned1 的任务主指标均值低于 baseline；n=1000 vs n=1000。")
    ws.cell(8, 10).font = Font(name="Aptos", size=9, italic=True, color=GREY)

    ws.cell(detail_end + 3, 1, f"生成时间：{generated_at}")
    ws.cell(detail_end + 3, 1).font = Font(name="Aptos", size=9, color=GREY)
    _set_widths(
        ws,
        {
            1: 35,
            2: 10,
            3: 23,
            4: 12,
            5: 15,
            6: 16,
            7: 16,
            8: 16,
            9: 3,
            10: 16,
            11: 16,
            12: 16,
            13: 16,
        },
    )
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A15"


def _write_test_sheet(wb: Workbook, test_type: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(test_type.upper())
    test_rows = [row for row in rows if row["test_type"] == test_type]
    _style_title(
        ws,
        f"{test_type.upper()}：调参后较优结果",
        "蓝色行采用 tuned1；其他行保留 baseline。Primary 指标随任务变化，定义见“说明”页。",
        15,
    )
    headers = [
        "PDE",
        "task",
        "sensor",
        "primary metric",
        "结果来源",
        "n",
        "最终 primary",
        "最终 rel L2(a)",
        "最终 rel L2(a) P90",
        "最终 rel L2(u)",
        "最终 rel L2(u) P90",
        "最终 residual",
        "Base primary",
        "Tuned1 primary",
        "主指标改善",
    ]
    header_row = 4
    for col_number, header in enumerate(headers, start=1):
        ws.cell(header_row, col_number, header)
    for row_number, row in enumerate(test_rows, start=header_row + 1):
        values = [
            row["pde"],
            row["task"],
            row["sensor_mode"],
            row["primary_metric"],
            row["result_source"],
            row["final_n"],
            row["final_primary_mean"],
            row["final_rel_l2_a_mean"],
            row["final_rel_l2_a_p90"],
            row["final_rel_l2_u_mean"],
            row["final_rel_l2_u_p90"],
            row["final_pde_residual_mean"],
            row["baseline_primary_mean"],
            row["tuned1_primary_mean"],
            row["primary_improvement"] if row["replaced"] else None,
        ]
        for col_number, value in enumerate(values, start=1):
            ws.cell(row_number, col_number, value)
    end_row = header_row + len(test_rows)
    _apply_table_style(ws, header_row, end_row, len(headers), source_col=5)
    for row_number in range(header_row + 1, end_row + 1):
        for col_number in (7, 8, 9, 10, 11, 12, 13, 14):
            ws.cell(row_number, col_number).number_format = "0.000000"
        ws.cell(row_number, 15).number_format = "0.0%"
    ws.conditional_formatting.add(
        f"O{header_row + 1}:O{end_row}",
        ColorScaleRule(
            start_type="min",
            start_color=BLUE_LIGHT,
            end_type="max",
            end_color=BLUE_DARK,
        ),
    )
    _set_widths(
        ws,
        {
            1: 22,
            2: 11,
            3: 16,
            4: 19,
            5: 12,
            6: 8,
            7: 15,
            8: 16,
            9: 19,
            10: 16,
            11: 19,
            12: 16,
            13: 15,
            14: 16,
            15: 14,
        },
    )
    ws.print_title_rows = f"1:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _write_audit_sheet(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("替换审计")
    _style_title(
        ws,
        "Baseline 与 tuned1 替换审计",
        "完整保留替换前后值与选择原因；这里不对 residual 单独择优。",
        19,
    )
    fields = [
        ("TEST", "test_type"),
        ("PDE", "pde"),
        ("task", "task"),
        ("sensor", "sensor_mode"),
        ("primary metric", "primary_metric"),
        ("结果来源", "result_source"),
        ("是否替换", "replaced"),
        ("选择原因", "selection_reason"),
        ("Base n", "baseline_n"),
        ("Tuned1 n", "tuned1_n"),
        ("Base primary", "baseline_primary_mean"),
        ("Tuned1 primary", "tuned1_primary_mean"),
        ("最终 primary", "final_primary_mean"),
        ("主指标改善", "primary_improvement"),
        ("Base residual", "baseline_pde_residual_mean"),
        ("Tuned1 residual", "tuned1_pde_residual_mean"),
        ("最终 residual", "final_pde_residual_mean"),
        ("Residual 改善", "pde_residual_improvement"),
        ("最终 n", "final_n"),
    ]
    header_row = 4
    for col_number, (label, _) in enumerate(fields, start=1):
        ws.cell(header_row, col_number, label)
    for row_number, row in enumerate(rows, start=header_row + 1):
        for col_number, (_, field) in enumerate(fields, start=1):
            value = row[field]
            if field == "test_type":
                value = value.upper()
            elif field == "replaced":
                value = "是" if value else "否"
            ws.cell(row_number, col_number, value)
    end_row = header_row + len(rows)
    _apply_table_style(ws, header_row, end_row, len(fields), source_col=6)
    for row_number in range(header_row + 1, end_row + 1):
        for col_number in (11, 12, 13, 15, 16, 17):
            ws.cell(row_number, col_number).number_format = "0.000000"
        for col_number in (14, 18):
            ws.cell(row_number, col_number).number_format = "0.0%"
    widths = {index: 15 for index in range(1, len(fields) + 1)}
    widths.update({2: 22, 5: 19, 8: 24})
    _set_widths(ws, widths)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _write_notes_sheet(wb: Workbook, source_path: Path, generated_at: str) -> None:
    ws = wb.create_sheet("说明")
    _style_title(ws, "口径与限制说明", "本页记录结果替换规则、指标定义和来源。", 8)
    notes = [
        ("替换规则", "若同一 TEST/PDE/task/sensor 下 tuned1 的任务主指标均值严格低于 baseline，则整行采用 tuned1；否则保留 baseline。"),
        ("forward primary", "rel_l2_u_mean"),
        ("inverse primary", "rel_l2_a_mean"),
        ("both primary", "max(rel_l2_a_mean, rel_l2_u_mean)"),
        ("改善率", "(baseline primary - tuned1 primary) / baseline primary"),
        ("行级一致性", "采用 tuned1 时，最终 rel L2(a)、rel L2(u)、P90 和 PDE residual 全部来自同一 tuned1 配置；不按单项指标拼接。"),
        ("样本量", "所有纳入的 baseline/tuned1 目标组均为 n=1000；smooth 中 n=1–4 的 pilot 组未纳入。"),
        ("重要限制", "SMOOTH 的 darcy/helmholtz/nsnonbounded/poisson inverse baseline 使用的 zeta 参数与 ID/ROUGH 不一致，跨 TEST 比较改善幅度时需谨慎。"),
        ("Residual 提醒", "替换依据仅为任务主指标；18 个替换项中有 11 个 PDE residual 上升，已在摘要和审计页完整保留。"),
        ("主数据源", str(source_path)),
        ("原始聚合", "outputs/main/MAIN1000_100_TEST_{id,rough,smooth}{,_tuned1}/summary_all_grouped.csv"),
        ("生成时间", generated_at),
    ]
    header_row = 4
    ws.cell(header_row, 1, "项目")
    ws.cell(header_row, 2, "说明")
    for row_number, (label, note) in enumerate(notes, start=header_row + 1):
        ws.cell(row_number, 1, label)
        ws.cell(row_number, 2, note)
        ws.cell(row_number, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row_number].height = 34 if len(note) > 60 else 24
    _apply_table_style(ws, header_row, header_row + len(notes), 2)
    _set_widths(ws, {1: 22, 2: 105})


def build_workbook(source_path: Path, output_path: Path) -> None:
    rows = _select_best(_read_comparisons(source_path))
    _validate_rows(rows)
    generated_at = datetime.now().astimezone().isoformat(timespec="seconds")

    wb = Workbook()
    wb.properties.title = "MAIN1000 调参后较优结果汇总"
    wb.properties.subject = "Baseline rows replaced by tuned1 only when task-primary mean improves"
    wb.properties.creator = "FM4PDE analysis"
    wb.properties.description = "Source-aware result selection with replacement audit."
    _write_summary(wb, rows, generated_at)
    for test_type in TEST_ORDER:
        _write_test_sheet(wb, test_type, rows)
    _write_audit_sheet(wb, rows)
    _write_notes_sheet(wb, source_path, generated_at)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    check = load_workbook(output_path, data_only=False, read_only=False)
    assert check.sheetnames == ["摘要", "ID", "ROUGH", "SMOOTH", "替换审计", "说明"]
    assert check["ID"].max_row == 18
    assert check["ROUGH"].max_row == 18
    assert check["SMOOTH"].max_row == 18
    assert check["替换审计"].max_row == 46
    assert len(check["摘要"]._charts) == 1
    check.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input",
        type=Path,
        default=Path(
            "outputs/sampling_results_analysis/main1000_tuned1_summary/comparison.csv"
        ),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(
            "outputs/sampling_results_analysis/main1000_tuned1_summary/"
            "MAIN1000_best_after_tuned1.xlsx"
        ),
    )
    args = parser.parse_args()
    build_workbook(args.input.resolve(), args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
