#!/usr/bin/env python3
"""Build concise Excel summaries for main and ablation experiments."""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import statistics
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from sampling.aggregate import collect_ablation_report_rows  # noqa: E402


METRIC_COLUMNS = (
    ("rel L2(a)", "rel_l2_a"),
    ("rel L2(u)", "rel_l2_u"),
    ("pde L", "L_pde"),
    ("obs L(a)", "L_obs_a"),
    ("obs L(u)", "L_obs_u"),
)
METRIC_OUTPUT_COLUMNS = tuple(
    column
    for label, _ in METRIC_COLUMNS
    for column in (label, f"{label} std", f"{label} n")
)
MAIN_COLUMNS = (
    "PDE",
    "TASK",
    "DIST",
    "SENSOR",
    *METRIC_OUTPUT_COLUMNS,
    "Remark",
)
MAIN_DIR_PATTERN = re.compile(
    r"^MAIN1000_100_TEST_(?:(FULL)_)?(id|rough|smooth)(?:_.+)?$",
    re.IGNORECASE,
)
DIST_NAMES = {"id": "ID", "rough": "Rough", "smooth": "Smooth"}
TASK_ORDER = {"both": 0, "forward": 1, "inverse": 2}

ABLATION_DIMENSIONS: dict[str, tuple[tuple[str, str], ...]] = {
    "guidance_components": (("GUIDANCE", "guidance_components"),),
    "loss_state_by_sampler": (
        ("LOSS STATE", "loss_state"),
        ("SAMPLER", "sampler_phase"),
    ),
    "noise_robustness": (("NOISE", "noise_level"),),
    "num_steps_by_sampler": (
        ("NUM STEPS", "num_steps"),
        ("SAMPLER", "sampler_phase"),
    ),
    "sampler_phase": (
        ("SAMPLER", "sampler_phase"),
        ("SWITCH RATIO", "switch_ratio"),
    ),
    "poisson_sampler_comparison": (
        ("SAMPLER", "sampler_phase"),
        ("SWITCH RATIO", "switch_ratio"),
    ),
    "sensor_mode": (("SENSOR", "sensor_mode"),),
    "sensor_sparsity": (("NUM OBS", "num_obs"),),
    "statistics_stability": (("NUM SEEDS", "rel_l2_a_n"),),
    "step_method_by_sampler": (
        ("STEP METHOD", "step_method"),
        ("SAMPLER", "sampler_phase"),
    ),
    "temporal_residual_mode": (
        ("RESIDUAL MODE", "residual_mode"),
        ("RESOLVED MODE", "resolved_residual_mode"),
    ),
    "time_grid_by_sampler": (
        ("TIME GRID", "time_grid"),
        ("SAMPLER", "sampler_phase"),
    ),
    "deterministic_endpoint_bt": (
        ("ENDPOINT MODE", "deterministic_endpoint_mode"),
        ("BT MODE", "deterministic_bt_mode"),
        ("GUIDANCE COEFF", "deterministic_guidance_coeff"),
        ("BT MAX SCALE", "deterministic_bt_max_scale"),
    ),
}

HEADER_FILL = PatternFill("solid", fgColor="17365D")
HEADER_FONT = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Aptos", size=10, color="172B4D")
ALT_FILL = PatternFill("solid", fgColor="F4F7FB")
THIN_GREY = Side(style="thin", color="D9E2F3")


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _number(value: Any) -> float | int | None:
    if value in (None, ""):
        return None
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(parsed):
        return None
    return int(parsed) if parsed.is_integer() else parsed


def _first_aggregate_stats(
    candidates: Sequence[tuple[Mapping[str, Any], str]],
) -> tuple[float | int | None, float | int | None, int | None]:
    """Return mean, sample standard deviation, and count from one metric source.

    Choosing all three fields from the same row and metric stem prevents, for
    example, pairing an ``L_pde`` mean with a ``pde_residual_norm`` standard
    deviation when the preferred aggregate is unavailable.
    """
    for row, metric in candidates:
        mean = _number(row.get(f"{metric}_mean"))
        if mean is None:
            continue
        std = _number(row.get(f"{metric}_std"))
        count = _number(row.get(f"{metric}_n"))
        return mean, std, int(count) if count is not None else None
    return None, None, None


def _set_metric_stats(
    result: dict[str, Any],
    label: str,
    stats: tuple[Any, Any, Any],
) -> None:
    mean, std, count = stats
    result[label] = mean
    result[f"{label} std"] = std
    result[f"{label} n"] = count


def _set_aggregate_metric_stats(
    result: dict[str, Any],
    sample_row: Mapping[str, Any],
    run_row: Mapping[str, Any],
) -> None:
    """Populate the five output metrics from matching sample/run aggregates."""
    _set_metric_stats(
        result,
        "rel L2(a)",
        _first_aggregate_stats(((sample_row, "rel_l2_a"),)),
    )
    _set_metric_stats(
        result,
        "rel L2(u)",
        _first_aggregate_stats(((sample_row, "rel_l2_u"),)),
    )
    _set_metric_stats(
        result,
        "pde L",
        _first_aggregate_stats(
            (
                (run_row, "L_pde"),
                (run_row, "pde_residual_norm"),
                (sample_row, "L_pde"),
                (sample_row, "pde_residual_norm"),
            )
        ),
    )
    _set_metric_stats(
        result,
        "obs L(a)",
        _first_aggregate_stats(
            (
                (run_row, "L_obs_a"),
                (run_row, "clean_L_obs_a"),
                (run_row, "obs_rel_l2_a"),
                (sample_row, "L_obs_a"),
                (sample_row, "clean_L_obs_a"),
                (sample_row, "obs_rel_l2_a"),
            )
        ),
    )
    _set_metric_stats(
        result,
        "obs L(u)",
        _first_aggregate_stats(
            (
                (run_row, "L_obs_u"),
                (run_row, "clean_L_obs_u"),
                (run_row, "obs_rel_l2_u"),
                (sample_row, "L_obs_u"),
                (sample_row, "clean_L_obs_u"),
                (sample_row, "obs_rel_l2_u"),
            )
        ),
    )


def _single_run_metric_stats(
    source: Mapping[str, Any], run_dir: Path
) -> dict[str, tuple[float | int | None, float | int | None, int | None]]:
    """Compute per-run statistics from per-sample metrics when available.

    Some losses are currently persisted only as one run-level aggregate. Such
    metrics report ``n=1`` and ``std=0`` instead of implying that a per-sample
    deviation was available.
    """
    sample_values: dict[str, list[float]] = defaultdict(list)
    sample_path = run_dir / "metrics_per_sample.csv"
    if sample_path.is_file():
        for row in _read_csv(sample_path):
            for _, metric in METRIC_COLUMNS:
                value = _number(row.get(metric))
                if value is not None:
                    sample_values[metric].append(float(value))

    stats: dict[str, tuple[float | int | None, float | int | None, int | None]] = {}
    for _, metric in METRIC_COLUMNS:
        mean = _number(source.get(metric))
        values = sample_values.get(metric, [])
        if values:
            std = statistics.stdev(values) if len(values) > 1 else 0.0
            stats[metric] = (mean, std, len(values))
        elif mean is not None:
            stats[metric] = (mean, 0.0, 1)
        else:
            stats[metric] = (None, None, None)
    return stats


def _normalize_sensor(value: Any) -> str:
    sensor = str(value or "")
    if sensor in {"sensor_column", "sensor_col", "sersor_col"}:
        return "sensor_col"
    return sensor


def _full_main_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    """Exclude partial pilot groups when a MAIN1000 count column is available."""
    if not rows or "rel_l2_a_n" not in rows[0]:
        return rows
    return [row for row in rows if (_number(row.get("rel_l2_a_n")) or 0) >= 1000]


def _main_group_key(row: Mapping[str, Any]) -> str:
    """Match sample/run aggregates despite derived residual-mode enrichment."""
    key = str(row.get("ablation_group_key", ""))
    return re.sub(
        r"(?:(?<=^)|(?<=\|))resolved_residual_mode=[^|]*",
        "resolved_residual_mode=",
        key,
    )


def collect_main_results(outputs_root: str | Path) -> dict[str, list[dict[str, Any]]]:
    """Collect all complete MAIN1000 result groups under ``outputs/main``."""
    outputs_root = Path(outputs_root)
    main_root = outputs_root / "main"
    if not main_root.is_dir():
        raise FileNotFoundError(f"Main experiment directory does not exist: {main_root}")

    collected: dict[str, list[dict[str, Any]]] = {"sparse": [], "full": []}
    matched_directories = 0
    for experiment_dir in sorted(path for path in main_root.iterdir() if path.is_dir()):
        match = MAIN_DIR_PATTERN.fullmatch(experiment_dir.name)
        if match is None:
            continue
        matched_directories += 1
        source_path = experiment_dir / "summary_all_grouped.csv"
        if not source_path.is_file():
            raise FileNotFoundError(f"Missing main result aggregate: {source_path}")

        loss_path = experiment_dir / "summary_run_seed_grouped.csv"
        loss_rows = _read_csv(loss_path) if loss_path.is_file() else []
        loss_by_key = {
            _main_group_key(row): row
            for row in loss_rows
            if row.get("ablation_group_key")
        }
        observation = "full" if match.group(1) else "sparse"
        distance = DIST_NAMES[match.group(2).lower()]

        for source_row in _full_main_rows(_read_csv(source_path)):
            loss_row = loss_by_key.get(_main_group_key(source_row), {})
            result = {
                "PDE": source_row.get("pde", ""),
                "TASK": source_row.get("task", ""),
                "DIST": distance,
                "SENSOR": _normalize_sensor(source_row.get("sensor_mode")),
                "Remark": experiment_dir.name,
            }
            _set_aggregate_metric_stats(result, source_row, loss_row)
            collected[observation].append(result)

    if matched_directories == 0:
        raise FileNotFoundError(f"No MAIN1000_100_TEST_* directories found in {main_root}")
    for rows in collected.values():
        rows.sort(
            key=lambda row: (
                str(row["PDE"]),
                TASK_ORDER.get(str(row["TASK"]), 99),
                tuple(DIST_NAMES.values()).index(str(row["DIST"])),
                str(row["SENSOR"]),
                str(row["Remark"]),
            )
        )
    return collected


def _excel_value(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    numeric = _number(value)
    return value if numeric is None else numeric


def _write_table(ws: Any, columns: Sequence[str], rows: Iterable[Mapping[str, Any]]) -> None:
    rows = list(rows)
    ws.append(list(columns))
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GREY)
    ws.row_dimensions[1].height = 28

    for row_number, row in enumerate(rows, start=2):
        ws.append([_excel_value(row.get(column)) for column in columns])
        for cell in ws[row_number]:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(bottom=THIN_GREY)
            if row_number % 2 == 0:
                cell.fill = ALT_FILL
        for column_number, column in enumerate(columns, start=1):
            if column in {
                name
                for label, _ in METRIC_COLUMNS
                for name in (label, f"{label} std")
            }:
                ws.cell(row_number, column_number).number_format = "0.000000E+00"
            elif column in {f"{label} n" for label, _ in METRIC_COLUMNS}:
                ws.cell(row_number, column_number).number_format = "0"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}"
    ws.sheet_view.showGridLines = False
    for column_number, column in enumerate(columns, start=1):
        values = [len(str(row.get(column, ""))) for row in rows]
        width = min(max([len(column), *values], default=len(column)) + 2, 48)
        if column == "Remark":
            width = min(max(width, 34), 64)
        ws.column_dimensions[get_column_letter(column_number)].width = width


def build_main_workbook(outputs_root: str | Path, output_path: str | Path) -> Path:
    results = collect_main_results(outputs_root)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, result_key in (
        ("FM4PDE_SPARSE", "sparse"),
        ("FM4PDE_FULL", "full"),
    ):
        worksheet = workbook.create_sheet(sheet_name)
        _write_table(worksheet, MAIN_COLUMNS, results[result_key])
    workbook.properties.title = "FM4PDE main experiment summary"
    workbook.properties.creator = "FM4PDE"
    workbook.save(output_path)
    return output_path


def _sheet_name(group: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", group or "ungrouped")[:31] or "ungrouped"
    name = base
    counter = 2
    while name in used:
        suffix = f"_{counter}"
        name = f"{base[: 31 - len(suffix)]}{suffix}"
        counter += 1
    used.add(name)
    return name


def _sort_value(value: Any) -> tuple[int, Any]:
    numeric = _number(value)
    return (0, numeric) if numeric is not None else (1, str(value or ""))


def collect_ablation_results(outputs_root: str | Path) -> dict[str, list[dict[str, Any]]]:
    """Collect latest grouped ablation statistics by experiment type."""
    outputs_root = Path(outputs_root).resolve()
    ablations_root = outputs_root / "ablations"
    if not ablations_root.is_dir():
        raise FileNotFoundError(f"Ablation experiment directory does not exist: {ablations_root}")

    sample_path = ablations_root / "summary_latest_grouped.csv"
    run_path = ablations_root / "summary_latest_run_seed_grouped.csv"
    report_path = ablations_root / "ablation_report_metrics.csv"
    sources: Iterable[Mapping[str, Any]]
    run_by_key: dict[str, Mapping[str, Any]] = {}
    using_aggregates = sample_path.is_file() and run_path.is_file()
    if using_aggregates:
        sources = _read_csv(sample_path)
        run_by_key = {
            _main_group_key(row): row
            for row in _read_csv(run_path)
            if row.get("ablation_group_key")
        }
    elif report_path.is_file():
        sources = _read_csv(report_path)
    else:
        sources = collect_ablation_report_rows(ablations_root)

    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for source in sources:
        group = str(source.get("ablation_group") or "ungrouped")
        dimensions = ABLATION_DIMENSIONS.get(group, (("NAME", "ablation_name"),))
        if using_aggregates:
            run_dir = Path()
            remark = "ablations/summary_latest_grouped.csv"
        else:
            run_dir = Path(str(source.get("run_dir", "")))
            try:
                remark = str(run_dir.resolve().relative_to(outputs_root))
            except (OSError, ValueError):
                remark = str(run_dir)
        result: dict[str, Any] = {
            "PDE": source.get("pde", ""),
            "TASK": source.get("task", ""),
            "DIST": DIST_NAMES.get(
                str(source.get("test_type", "")).lower(),
                source.get("test_type", ""),
            ),
            "Remark": remark,
        }
        if using_aggregates:
            run_row = run_by_key.get(_main_group_key(source), {})
            _set_aggregate_metric_stats(result, source, run_row)
        else:
            run_stats = _single_run_metric_stats(source, run_dir)
            for label, metric in METRIC_COLUMNS:
                _set_metric_stats(result, label, run_stats[metric])
        for label, field in dimensions:
            value = source.get(field, "")
            if value == "" and field == "rel_l2_a_n":
                value = result.get("rel L2(a) n", "")
            result[label] = value
        grouped[group].append(result)

    if not grouped:
        raise FileNotFoundError(f"No analysis-ready ablation runs found in {ablations_root}")
    for group, rows in grouped.items():
        dimensions = ABLATION_DIMENSIONS.get(group, (("NAME", "ablation_name"),))
        labels = [label for label, _ in dimensions]
        rows.sort(
            key=lambda row: (
                str(row["PDE"]),
                TASK_ORDER.get(str(row["TASK"]), 99),
                str(row["DIST"]),
                *(_sort_value(row.get(label)) for label in labels),
                str(row["Remark"]),
            )
        )
    return dict(sorted(grouped.items()))


def build_ablation_workbook(outputs_root: str | Path, output_path: str | Path) -> Path:
    grouped = collect_ablation_results(outputs_root)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    workbook.remove(workbook.active)
    used_sheet_names: set[str] = set()
    for group, rows in grouped.items():
        dimensions = ABLATION_DIMENSIONS.get(group, (("NAME", "ablation_name"),))
        columns = (
            "PDE",
            "TASK",
            "DIST",
            *(label for label, _ in dimensions),
            *METRIC_OUTPUT_COLUMNS,
            "Remark",
        )
        worksheet = workbook.create_sheet(_sheet_name(group, used_sheet_names))
        _write_table(worksheet, columns, rows)
    workbook.properties.title = "FM4PDE ablation experiment summary"
    workbook.properties.creator = "FM4PDE"
    workbook.save(output_path)
    return output_path


def summary(outputs_root: str | Path = "outputs", exp: str = "all") -> dict[str, Path]:
    """Write the selected experiment summaries into ``outputs/summary``."""
    if exp not in {"all", "main", "ablations"}:
        raise ValueError("exp must be one of: all, main, ablations")
    outputs_root = Path(outputs_root).resolve()
    summary_root = outputs_root / "summary"
    summary_root.mkdir(parents=True, exist_ok=True)

    written: dict[str, Path] = {}
    if exp in {"all", "main"}:
        written["main"] = build_main_workbook(
            outputs_root, summary_root / "main_summary.xlsx"
        )
    if exp in {"all", "ablations"}:
        written["ablations"] = build_ablation_workbook(
            outputs_root, summary_root / "ablations_summary.xlsx"
        )
    return written


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--exp",
        choices=("all", "main", "ablations"),
        default="all",
        help="Experiment family to summarize (default: all).",
    )
    parser.add_argument(
        "--outputs-root",
        type=Path,
        default=Path("outputs"),
        help="Outputs root containing main/ and ablations/ (default: outputs).",
    )
    return parser


def main(argv: list[str] | None = None) -> dict[str, Path]:
    args = build_arg_parser().parse_args(argv)
    written = summary(args.outputs_root, args.exp)
    for experiment, path in written.items():
        print(f"{experiment}: {path}")
    return written


if __name__ == "__main__":
    main()
